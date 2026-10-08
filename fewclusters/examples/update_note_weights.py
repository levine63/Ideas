"""Add the authorized weighting discussion and Appendix A to the supplied note."""
import argparse
import csv
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ap=argparse.ArgumentParser(); ap.add_argument('source'); ap.add_argument('output'); a=ap.parse_args()
d=Document(a.source)
root=Path(__file__).resolve().parents[1]
original=list(d.paragraphs)
anchor=next(p for p in original if p.text.startswith('4. How accurate'))
def insert(text,style='Body Text'):
    return anchor.insert_paragraph_before(text,style)
def link(p,text,anchor_name=None,url=None):
    h=OxmlElement('w:hyperlink')
    if anchor_name:h.set(qn('w:anchor'),anchor_name)
    if url:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

status=next(p for p in original if p.text=='Abstract').insert_paragraph_before('PRELIMINARY - NOT FOR PUBLICATION','Normal')
status.runs[0].bold=True
h=insert('3.1 Choosing weights across clusters','Heading 3')
insert('Under a common-effect model with independent unbiased cluster estimates and known variances, inverse-variance weighting minimizes variance among linear unbiased combinations. This oracle result does not establish that estimated precision weights are optimal for the sign test. Estimating variances from the same observations used to estimate effects can distort inference, especially with skewed outcomes.')
p=insert('In our preliminary simulations, estimated precision weights caused undercoverage under skewed errors, whereas size weighting improved power over equal score weights and maintained coverage near the nominal level across the designs considered. We therefore recommend size weighting as the primary weighting comparison for the preliminary analysis, retaining estimated precision weighting as experimental. ')
link(p,'Appendix A',anchor_name='appendix_weights')
p.add_run(' reports the designs, Monte Carlo uncertainty, and a diagnostic that separates nuisance fitting from same-sample variance weighting. These findings are not a general robustness theorem.')
insert('Here size weighting means weights proportional to cluster sample size on the cluster effect estimates. Because the scores in (3) already multiply each estimate by the square root of cluster sample size, this is implemented by replacing the equal score weights in (4) with weights proportional to that square root. Equal score weights already give square-root-of-size weights on cluster estimates. With heterogeneous treatment effects, changing these weights can change the estimand. Downweighting small clusters does not establish normality of their scores or extend the assumptions of Proposition 1.')
if hasattr(d,'add_comment'):
    d.add_comment(h.runs,text='Added the requested qualified weighting recommendation and a cross-reference to the new preliminary simulation appendix. This distinction concerns weights across effect estimates, not the local/borrowed prediction weights.',author='Codex',initials='CX')

p=original[88]
p.add_run(' Appendix A supplies preliminary weighting stress tests; comparisons using the exact empirical learners and broader sampling designs remain to be completed.')
d.add_page_break()
h=d.add_paragraph('Appendix A Cluster weighting simulations','Heading 2')
b=OxmlElement('w:bookmarkStart'); b.set(qn('w:id'),'901');b.set(qn('w:name'),'appendix_weights');h._p.insert(0,b)
e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),'901');h._p.append(e)
d.add_paragraph('PRELIMINARY - NOT FOR PUBLICATION','Normal').runs[0].bold=True
d.add_paragraph('A1 Design and weighting rules','Heading 3')
d.add_paragraph('We use the 16 JTPA site sizes: 38, 74, 81, 87, 177, 179, 190, 234, 353, 401, 463, 485, 524, 636, 788, and 1,392 observations. Within sites, observations are independent and treatment is independently assigned with probability two thirds. The common treatment effect is 0.08 in simulation units, not JTPA dollars. Outcomes depend linearly on two independent standard-normal covariates, with site-specific intercepts and slopes, plus treatment and an independent error.','Body Text')
d.add_paragraph('Four designs use Gaussian errors with equal site variances; Gaussian errors with unequal site variances; standardized Student-t errors with three degrees of freedom and unequal site variances; and centered, population-standardized lognormal errors with log standard deviation 1.5 and unequal site variances. In the unequal-variance designs, the site error standard deviation is the fourth root of the ratio of median site size to site size. Thus small sites are also noisier per observation. The Student-t design has finite variance but no finite fourth moment. The lognormal design stresses skewness and tails; it is not an estimated earnings distribution.','Body Text')
d.add_paragraph('Each design has 1,000 replications with oracle nuisance functions and 500 with fitted nuisances. Fitted outcome models use five-fold adaptive local and pooled regressions with site indicators, Ridge penalty one, outcome-pooling shrinkage 20, and a 20-observation local-training floor. Propensity is either known or estimated using adaptively pooled local and pooled sample means. Shared training across sites is included. These simple learners isolate the weighting question; they do not validate the exact boosting learners used in the JTPA application.','Body Text')
d.add_paragraph('We compare equal score weights, size weights, and estimated precision weights. Precision weights use a within-site sample variance of estimated influence contributions. The variance is shrunk toward its pooled target with strength 100 and floored at one tenth of that target; strength 20 is a prespecified sensitivity. These variance-shrinkage settings differ from outcome-pooling shrinkage. All weights remain fixed over sign transformations and hypothesized effects. Coverage is acceptance of the true effect by the inverted 5% test; power is rejection of zero. We enumerate all sign pairs and retain every replication, with no fitting failures or exclusions.','Body Text')
d.add_paragraph('A2 Coverage and power','Heading 3')
d.add_paragraph('Table A1 reports coverage and power percentages from the 500 fitted-nuisance replications. Each cell is coverage / power. At a rejection probability of 5%, the Monte Carlo standard error is about one percentage point. The target coverage is 95%.','Body Text')
rows=list(csv.DictReader((root/'results/weighting_simulation/summary.csv').open()))
t=d.add_table(rows=1,cols=5);t.autofit=False
widths=[1.55,1.0,1.1,1.1,1.2]
for c,w in zip(t.columns,widths):c.width=Inches(w)
for cell,text in zip(t.rows[0].cells,['Errors','Propensity','Current','Size','Precision']):cell.text=text
labels={'normal_equal':'Gaussian equal variance','normal_unequal':'Gaussian unequal variance','t3_unequal':'Student t unequal variance','lognormal_unequal':'Lognormal unequal variance'}
for design in labels:
 for method,label in [('shrink known p','Known'),('shrink estimated p','Estimated')]:
    cells=t.add_row().cells;cells[0].text=labels[design];cells[1].text=label
    for cell,rule in zip(cells[2:],['current','size','precision k100']):
        r=next(x for x in rows if x['design']==design and x['method']==method and x['weighting']==rule)
        cell.text=f"{100*float(r['coverage']):.1f} / {100*float(r['power']):.1f}"
for idx,row in enumerate(t.rows):
    pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    if idx==0:pr.append(OxmlElement('w:tblHeader'))
    for cell,w in zip(row.cells,widths):
        cell.width=Inches(w)
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        tcpr.append(borders)
        if idx==0:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');tcpr.append(shade)
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
            for run in p.runs:run.font.size=Pt(10);run.bold=idx==0
            if cell.text[:1].isdigit():p.alignment=WD_ALIGN_PARAGRAPH.CENTER

d.add_paragraph('Under lognormal errors, primary precision weights yield 90.6% coverage with either fitted propensity specification; the marginal 95% Monte Carlo interval is 87.7–92.9%. Size weights yield 95.2% coverage, with interval 93.0–96.8%. The weaker variance-shrinkage sensitivity gives 90.0% and 90.6% coverage. High alternative rejection rates from a test that overrejects the true null are not a fair efficiency gain. Across the four designs, size weighting raises power relative to current weights, while its observed coverage ranges from 93.6% to 96.0%. These marginal Monte Carlo intervals are not adjusted for multiple comparisons.','Body Text')
d.add_paragraph('A3 Diagnostic and implications','Heading 3')
d.add_paragraph('The skewed-error failure remains with oracle nuisances: same-sample precision weights give 91.3% coverage and a mean estimate of 0.08913 when the true effect is 0.08. In a post hoc diagnostic using the same 1,000 evaluation samples, we estimate precision weights from an additional independent sample with the same 6,102 observations and site-size distribution. Coverage rises to 94.5% and the mean estimate becomes 0.07996. Fixed weights based on the DGP variances give 94.7% coverage. This isolates same-sample variance weighting as a contributor to the failure in this design, rather than nuisance fitting alone.','Body Text')
d.add_paragraph('Dependence between skewed estimation errors and their estimated variances can bias weighting and spoil joint sign symmetry. Holding estimated weights fixed during enumeration does not remove that dependence. The diagnostic uses extra data and is not an equal-data-budget comparison or a validated holdout remedy; its power is lower than simple size weighting here. The simulations assume a common effect and independent within-site observations. They do not establish robustness to heterogeneous effects, common shocks, or arbitrary within-site dependence, and downweighting does not by itself justify small-site Gaussian approximations.','Body Text')
p=d.add_paragraph('Reproducibility. The evaluation seed is 20261009; the independent-weight diagnostic seed is 20261010. The full report includes all weighting sensitivities, oracle controls, paired comparisons, Monte Carlo intervals, source hashes, and replication results. See ','Body Text')
link(p,'the archived coverage and power report',url='https://github.com/levine63/Ideas/blob/8a7da00/fewclusters/results/weighting_simulation/COVERAGE_POWER.md')
p.add_run('. The recommendation concerns this preliminary analysis; publication requires further validation.')
Path(a.output).parent.mkdir(parents=True,exist_ok=True)
d.save(a.output)
# Every source native equation must be preserved byte-for-byte.
from lxml import etree
before=[etree.tostring(x) for p in original for x in p._p.xpath('.//m:oMath')]
after=[etree.tostring(x) for p in Document(a.output).paragraphs for x in p._p.xpath('.//m:oMath')]
assert before==after, 'Source equations changed'
print('Saved',a.output,'; preserved',len(before),'native equations')
